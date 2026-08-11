from __future__ import annotations

import re
from collections import defaultdict
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any, cast

from django.db import transaction
from django.db.models import Prefetch, Q, QuerySet
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from apps.articles.models import Article

from .models import RepostRecord

INVALID_SHEET_CHARACTERS = re.compile(r"[:\\/?*\[\]]")
AVAILABILITY_LABELS = {
    "UNKNOWN": "未知",
    "AVAILABLE": "正常",
    "UNAVAILABLE": "不可访问",
    "REMOVED": "已删除",
}


@dataclass(frozen=True)
class GlobalExportFilters:
    start_date: date | None = None
    end_date: date | None = None
    article_id: int | None = None
    author: str = ""
    monitoring_status: str = ""
    site_domain: str = ""
    has_repost: bool | None = None


def safe_excel_text(value: object) -> str:
    text = "" if value is None else str(value)
    return f"'{text}" if text.lstrip().startswith(("=", "+", "-", "@")) else text


def excel_datetime(value: date | datetime | None) -> date | datetime | None:
    if isinstance(value, datetime) and timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def sanitize_excel_sheet_name(value: str, used_names: set[str]) -> str:
    base = INVALID_SHEET_CHARACTERS.sub("_", value).strip(" '") or "未命名网站"
    base = base[:31]
    candidate = base
    suffix_number = 2
    while candidate.casefold() in {name.casefold() for name in used_names}:
        suffix = f"_{suffix_number}"
        candidate = f"{base[: 31 - len(suffix)]}{suffix}"
        suffix_number += 1
    used_names.add(candidate)
    return candidate


def _historical_records(as_of: datetime) -> QuerySet[RepostRecord]:
    return (
        RepostRecord.objects.filter(is_valid=True, created_at__lte=as_of)
        .select_related("platform", "article")
        .order_by("first_found_at", "first_discovered_at", "id")
    )


def filtered_articles(filters: GlobalExportFilters, *, as_of: datetime) -> QuerySet[Article]:
    queryset = Article.objects.filter(Q(retention_until__isnull=True) | Q(retention_until__gte=as_of))
    if filters.start_date:
        queryset = queryset.filter(published_date__gte=filters.start_date)
    if filters.end_date:
        queryset = queryset.filter(published_date__lte=filters.end_date)
    if filters.article_id:
        queryset = queryset.filter(id=filters.article_id)
    if filters.author:
        queryset = queryset.filter(Q(author__icontains=filters.author) | Q(author_department__icontains=filters.author))
    if filters.monitoring_status:
        queryset = queryset.filter(monitoring_status=filters.monitoring_status)
    if filters.site_domain:
        queryset = queryset.filter(
            repost_records__site_domain__iexact=filters.site_domain, repost_records__is_valid=True
        )
    if filters.has_repost is True:
        queryset = queryset.filter(repost_records__is_valid=True)
    elif filters.has_repost is False:
        queryset = queryset.exclude(repost_records__is_valid=True)
    return queryset.distinct().order_by("published_at", "published_date", "id")


def _record_domain(record: RepostRecord) -> str:
    platform = record.platform
    return record.site_domain or (platform.code.lower() if platform is not None else "unknown")


def _record_site_name(record: RepostRecord) -> str:
    platform = record.platform
    return record.site_name or (platform.name if platform is not None else _record_domain(record))


def _record_url(record: RepostRecord) -> str:
    return record.canonical_url or record.normalized_url or record.final_url or record.original_url


def _set_hyperlink(cell: Any, url: str, label: str | None = None) -> None:
    cell.value = safe_excel_text(label or url)
    cell.hyperlink = url
    cell.style = "Hyperlink"


def _format_sheet(sheet: Worksheet, widths: Iterable[int]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)


@transaction.atomic
def build_global_repost_workbook(filters: GlobalExportFilters) -> tuple[bytes, datetime]:
    """Build one read-only, point-in-time workbook without scheduling searches."""

    export_as_of = timezone.now()
    records_prefetch = Prefetch(
        "repost_records",
        queryset=_historical_records(export_as_of),
        to_attr="export_repost_records",
    )
    articles = list(filtered_articles(filters, as_of=export_as_of).prefetch_related(records_prefetch))
    records = [
        record
        for article in articles
        for record in cast(list[RepostRecord], getattr(article, "export_repost_records", []))
    ]
    if filters.site_domain:
        records = [record for record in records if _record_domain(record).casefold() == filters.site_domain.casefold()]

    records_by_article_site: dict[tuple[int, str], list[RepostRecord]] = defaultdict(list)
    site_names: dict[str, str] = {}
    for record in records:
        domain = _record_domain(record)
        site_names.setdefault(domain, _record_site_name(record))
        records_by_article_site[(record.article_id, domain)].append(record)
    domains = sorted(site_names, key=lambda item: (site_names[item].casefold(), item.casefold()))

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "数据"
    fixed_headers = [
        "发布时间",
        "作者",
        "标题",
        "原文章链接",
        "监测状态",
        "监测截止时间",
        "转载网站数",
        "转载链接数",
    ]
    data_sheet.append([*fixed_headers, *[site_names[domain] for domain in domains]])
    for article in articles:
        article_records = [record for record in records if record.article_id == article.id]
        distinct_domains = {_record_domain(record) for record in article_records}
        distinct_urls = {_record_url(record) for record in article_records}
        published = excel_datetime(article.published_at) or article.published_date
        data_sheet.append(
            [
                published,
                safe_excel_text(article.author or article.author_department),
                safe_excel_text(article.title),
                "",
                article.monitoring_status,
                excel_datetime(article.monitor_until),
                len(distinct_domains),
                len(distinct_urls),
            ]
        )
        row = data_sheet.max_row
        if article.original_url:
            _set_hyperlink(data_sheet.cell(row, 4), article.original_url)
        for offset, domain in enumerate(domains, start=len(fixed_headers) + 1):
            found = records_by_article_site.get((article.id, domain), [])
            if not found:
                continue
            first_url = _record_url(found[0])
            label = f"{first_url} √" + (f"（共{len(found)}条）" if len(found) > 1 else "")
            _set_hyperlink(data_sheet.cell(row, offset), first_url, label)
    _format_sheet(data_sheet, [19, 16, 48, 45, 14, 19, 14, 14, *([45] * len(domains))])

    used_names = {"数据", "总统计", "统计汇总"}
    for domain in domains:
        site_sheet = workbook.create_sheet(sanitize_excel_sheet_name(site_names[domain], used_names))
        site_sheet.append(
            [
                "文章名称",
                "转载链接",
                "原发布时间",
                "作者",
                "原文章链接",
                "转载标题",
                "标题相似度",
                "首次发现时间",
                "最后发现时间",
                "当前链接状态",
            ]
        )
        for record in [item for item in records if _record_domain(item) == domain]:
            article = record.article
            site_sheet.append(
                [
                    safe_excel_text(article.title),
                    "",
                    excel_datetime(article.published_at) or article.published_date,
                    safe_excel_text(article.author or article.author_department),
                    "",
                    safe_excel_text(record.result_title or record.repost_title),
                    float(record.similarity_score) if record.similarity_score is not None else "",
                    excel_datetime(record.first_found_at or record.first_discovered_at),
                    excel_datetime(record.last_seen_at or record.last_checked_at),
                    AVAILABILITY_LABELS.get(record.availability_status, "未知"),
                ]
            )
            row = site_sheet.max_row
            _set_hyperlink(site_sheet.cell(row, 2), _record_url(record))
            if article.original_url:
                _set_hyperlink(site_sheet.cell(row, 5), article.original_url)
        _format_sheet(site_sheet, [48, 48, 19, 16, 45, 48, 14, 19, 19, 14])

    total_sheet = workbook.create_sheet("总统计")
    total_sheet.append(["标题", "作者", "发布时间", "转载网站数", "转载链接数", *[site_names[d] for d in domains]])
    for article in articles:
        article_records = [record for record in records if record.article_id == article.id]
        found_domains = {_record_domain(record) for record in article_records}
        total_sheet.append(
            [
                safe_excel_text(article.title),
                safe_excel_text(article.author or article.author_department),
                excel_datetime(article.published_at) or article.published_date,
                len(found_domains),
                len({_record_url(record) for record in article_records}),
                *["√" if domain in found_domains else "" for domain in domains],
            ]
        )
    _format_sheet(total_sheet, [48, 16, 19, 14, 14, *([14] * len(domains))])

    summary_sheet = workbook.create_sheet("统计汇总")
    articles_with_reposts = {record.article_id for record in records}
    summary_sheet.append(["统计项", "数值", "导出时点"])
    for label, value in [
        ("原创文章总数", len(articles)),
        ("至少发现转载的文章数", len(articles_with_reposts)),
        ("暂无转载文章数", len(articles) - len(articles_with_reposts)),
        ("不同转载网站总数", len(domains)),
        ("转载链接总数", len({_record_url(record) for record in records})),
    ]:
        summary_sheet.append([label, value, excel_datetime(export_as_of)])
    summary_sheet.append([])
    summary_sheet.append(["转载网站", "转载域名", "曾转载原创文章数", "转载链接数", "已转载文章覆盖比例"])
    denominator = len(articles_with_reposts)
    for domain in domains:
        site_records = [record for record in records if _record_domain(record) == domain]
        site_article_count = len({record.article_id for record in site_records})
        summary_sheet.append(
            [
                safe_excel_text(site_names[domain]),
                safe_excel_text(domain),
                site_article_count,
                len({_record_url(record) for record in site_records}),
                site_article_count / denominator if denominator else 0,
            ]
        )
        summary_sheet.cell(summary_sheet.max_row, 5).number_format = "0.0%"
    _format_sheet(summary_sheet, [28, 32, 22, 16, 22])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), export_as_of
