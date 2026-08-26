from __future__ import annotations

from collections import defaultdict
from datetime import date
from hashlib import sha256
from io import BytesIO
from typing import Any

from django.core.files.base import ContentFile
from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.accounts.models import User
from apps.articles.models import Article
from apps.monitoring.models import PlatformDetectionStatus, StatusSnapshot
from apps.platforms.models import Platform
from apps.reposts.models import RepostRecord

from .models import GeneratedReport

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(name="Carlito", size=11, bold=True, color="000000")
BODY_FONT = Font(name="Carlito", size=11, color="000000")


def safe_excel_text(value: object) -> str:
    text = "" if value is None else str(value)
    return f"'{text}" if text.lstrip().startswith(("=", "+", "-", "@")) else text


def _status_symbol(status: str) -> str:
    symbols: dict[str, str] = {
        str(PlatformDetectionStatus.FOUND): "√",
        str(PlatformDetectionStatus.NOT_FOUND): "×",
        str(PlatformDetectionStatus.UNKNOWN): "—",
    }
    return symbols.get(status, "—")


def _status_display(status: str) -> str:
    return "0" if status == PlatformDetectionStatus.NOT_FOUND else "—"


def report_dataset(
    *, snapshot: StatusSnapshot, period_start: date, period_end: date
) -> dict[str, Any]:
    matrix = snapshot.matrix_data
    batch = snapshot.source_batch
    article_ids = (
        list(batch.article_ids)
        if batch
        else sorted({int(item["article_id"]) for item in matrix.values()})
    )
    platform_ids = (
        list(batch.platform_ids)
        if batch
        else sorted({int(item["platform_id"]) for item in matrix.values()})
    )
    articles = list(
        Article.objects.filter(
            id__in=article_ids, published_date__range=(period_start, period_end)
        ).order_by("published_date", "id")
    )
    platforms = list(
        Platform.objects.filter(id__in=platform_ids).order_by("name", "id")
    )
    repost_rows = list(
        RepostRecord.objects.filter(
            article_id__in=article_ids,
            platform_id__in=platform_ids,
            first_discovered_at__lte=snapshot.cutoff_at,
            is_valid=True,
        ).order_by("first_discovered_at", "id")
    )
    reposts_by_pair: dict[tuple[int, int], list[RepostRecord]] = defaultdict(list)
    for repost in repost_rows:
        if repost.platform_id is None:
            continue
        reposts_by_pair[(repost.article_id, repost.platform_id)].append(repost)
    rows: list[dict[str, Any]] = []
    for article in articles:
        cells: dict[int, dict[str, Any]] = {}
        found_count = completed_count = 0
        for platform in platforms:
            item = matrix.get(f"{article.id}:{platform.id}", {})
            status = str(item.get("status", PlatformDetectionStatus.UNKNOWN))
            records = reposts_by_pair[(article.id, platform.id)]
            if records:
                status = str(PlatformDetectionStatus.FOUND)
            if status == PlatformDetectionStatus.FOUND:
                found_count += 1
            if status in {
                PlatformDetectionStatus.FOUND,
                PlatformDetectionStatus.NOT_FOUND,
            }:
                completed_count += 1
            cells[platform.id] = {"status": status, "records": records}
        total = len(platforms)
        rows.append(
            {
                "article": article,
                "cells": cells,
                "repost_rate": found_count / total if total else 0,
                "completion_rate": completed_count / total if total else 0,
            }
        )
    platform_statistics: dict[int, dict[str, float]] = {}
    for platform in platforms:
        count = len(rows)
        found = sum(
            row["cells"][platform.id]["status"] == PlatformDetectionStatus.FOUND
            for row in rows
        )
        completed = sum(
            row["cells"][platform.id]["status"]
            in {PlatformDetectionStatus.FOUND, PlatformDetectionStatus.NOT_FOUND}
            for row in rows
        )
        platform_statistics[platform.id] = {
            "repost_rate": found / count if count else 0,
            "completion_rate": completed / count if count else 0,
        }
    new_reposts = list(
        RepostRecord.objects.filter(
            platform_id__in=platform_ids,
            first_discovered_at__date__range=(period_start, period_end),
            is_valid=True,
        )
        .select_related("article", "platform")
        .order_by("first_discovered_at", "id")
    )
    return {
        "articles": articles,
        "platforms": platforms,
        "rows": rows,
        "platform_statistics": platform_statistics,
        "repost_rows": repost_rows,
        "new_reposts": new_reposts,
    }


def _style_header(sheet: Any, cell_range: str) -> None:
    for row in sheet[cell_range]:
        for cell in row:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")


def _set_link(cell: Any, url: str) -> None:
    cell.value = safe_excel_text(url)
    cell.hyperlink = url
    cell.style = "Hyperlink"


def build_report_workbook(
    *, snapshot: StatusSnapshot, period_start: date, period_end: date
) -> tuple[bytes, dict[str, Any]]:
    data = report_dataset(
        snapshot=snapshot, period_start=period_start, period_end=period_end
    )
    platforms: list[Platform] = data["platforms"]
    rows: list[dict[str, Any]] = data["rows"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    headers = [
        "发布时间",
        "标题",
        "作者/部门",
        "总转载率",
        *[platform.name for platform in platforms],
        "检测完成率",
    ]
    sheet.append(headers)
    _style_header(sheet, f"A1:{chr(65 + len(headers) - 1)}1")
    for row in rows:
        article = row["article"]
        values: list[object] = [
            article.published_date,
            safe_excel_text(article.title),
            safe_excel_text(article.author_department),
            row["repost_rate"],
        ]
        values.extend("" for _ in platforms)
        values.append(row["completion_rate"])
        sheet.append(values)
        row_number = sheet.max_row
        for index, platform in enumerate(platforms, start=5):
            item = row["cells"][platform.id]
            cell = sheet.cell(row_number, index)
            if item["status"] == PlatformDetectionStatus.FOUND and item["records"]:
                _set_link(
                    cell,
                    item["records"][0].final_url or item["records"][0].original_url,
                )
            else:
                cell.value = _status_display(item["status"])
    for item in sheet["A"]:
        if item.row > 1:
            item.number_format = "yyyy-mm-dd"
    for column in (4, len(headers)):
        for item in sheet.iter_cols(min_col=column, max_col=column, min_row=2):
            for cell in item:
                cell.number_format = "0.0%"
    widths = [13, 46, 18, 12, *([34] * len(platforms)), 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    for platform in platforms:
        platform_sheet = workbook.create_sheet(platform.name[:31])
        platform_sheet.append(["文章名称", "链接", "转载发布时间", "首次发现时间"])
        _style_header(platform_sheet, "A1:D1")
        for repost in [
            item for item in data["repost_rows"] if item.platform_id == platform.id
        ]:
            platform_sheet.append(
                [
                    safe_excel_text(repost.article.title),
                    "",
                    (
                        repost.repost_published_at.strftime("%Y-%m-%d %H:%M:%S")
                        if repost.repost_published_at
                        else "无"
                    ),
                    repost.first_discovered_at.strftime("%Y-%m-%d %H:%M:%S"),
                ]
            )
            _set_link(
                platform_sheet.cell(platform_sheet.max_row, 2),
                repost.final_url or repost.original_url,
            )
        for column_name, width in {"A": 48, "B": 50, "C": 18, "D": 19}.items():
            platform_sheet.column_dimensions[column_name].width = width

    new_links = workbook.create_sheet("新增转载链接")
    new_links.append(["文章名称", "平台", "链接", "转载发布时间", "首次发现时间"])
    _style_header(new_links, "A1:E1")
    for repost in data["new_reposts"]:
        new_links.append(
            [
                safe_excel_text(repost.article.title),
                safe_excel_text(repost.platform.name),
                "",
                (
                    repost.repost_published_at.strftime("%Y-%m-%d %H:%M:%S")
                    if repost.repost_published_at
                    else "无"
                ),
                repost.first_discovered_at.strftime("%Y-%m-%d %H:%M:%S"),
            ]
        )
        _set_link(
            new_links.cell(new_links.max_row, 3),
            repost.final_url or repost.original_url,
        )
    for column_name, width in {"A": 48, "B": 20, "C": 50, "D": 18, "E": 19}.items():
        new_links.column_dimensions[column_name].width = width

    summary = workbook.create_sheet("总统计")
    summary.append(
        ["标题", "转载率", *[platform.name for platform in platforms], "检测完成率"]
    )
    _style_header(summary, f"A1:{chr(65 + len(platforms) + 2)}1")
    for row in rows:
        summary.append(
            [
                safe_excel_text(row["article"].title),
                row["repost_rate"],
                *[
                    _status_symbol(row["cells"][platform.id]["status"])
                    for platform in platforms
                ],
                row["completion_rate"],
            ]
        )
    summary.append([])
    summary.append(["平台统计", "", *[platform.name for platform in platforms]])
    summary.append(
        [
            "平台转载率",
            "",
            *[
                data["platform_statistics"][platform.id]["repost_rate"]
                for platform in platforms
            ],
        ]
    )
    summary.append(
        [
            "检测完成率",
            "",
            *[
                data["platform_statistics"][platform.id]["completion_rate"]
                for platform in platforms
            ],
        ]
    )
    _style_header(
        summary, f"A{len(rows) + 3}:{chr(65 + len(platforms) + 1)}{len(rows) + 3}"
    )
    for cell in summary["B"]:
        if cell.row > 1:
            cell.number_format = "0.0%"
    for row_index in (len(rows) + 4, len(rows) + 5):
        for cell in summary[row_index][2:]:
            cell.number_format = "0.0%"
    summary.column_dimensions["A"].width = 48
    summary.column_dimensions["B"].width = 12
    for index in range(3, len(platforms) + 4):
        summary.column_dimensions[chr(64 + index)].width = 14

    notes = workbook.create_sheet("说明")
    notes.append(["项目", "说明"])
    _style_header(notes, "A1:B1")
    for pair in [
        ("报表样式", "简单表格，不包含图表、仪表盘、KPI 卡片、趋势图或可视化大屏。"),
        (
            "状态1",
            "发现转载时显示真实 Excel 超链接；同平台多条链接保存在对应平台工作表。",
        ),
        ("状态0", "检测成功但未发现转载，数据表显示 0，总统计显示 ×。"),
        ("状态—", "未检测、超时、验证码、登录限制、访问失败或解析失败，总统计显示 —。"),
        ("转载率", "已转载平台数 ÷ 本批次选择平台总数。"),
        ("检测完成率", "状态为 1 或 0 的平台数 ÷ 本批次选择平台总数。"),
        ("新转载链接", f"本统计范围内首次发现链接数：{len(data['new_reposts'])}。"),
        ("阅读量", "V1.0 暂不采集，不放入主报表。"),
    ]:
        notes.append([safe_excel_text(pair[0]), safe_excel_text(pair[1])])
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 92
    for work_sheet in workbook.worksheets:
        for row in work_sheet.iter_rows():
            for cell in row:
                if cell.row != 1 and cell.style != "Hyperlink":
                    cell.font = BODY_FONT
                cell.alignment = Alignment(
                    vertical="center", wrap_text=work_sheet.title == "说明"
                )
        work_sheet.freeze_panes = "A2"

    output = BytesIO()
    workbook.save(output)
    statistics = {
        "period_start": period_start.isoformat(),
        "period_end": period_end.isoformat(),
        "article_total": len(rows),
        "platform_total": len(platforms),
        "new_repost_link_total": len(data["new_reposts"]),
    }
    return output.getvalue(), statistics


def create_report(
    *,
    report_type: str,
    report_date: date,
    period_start: date,
    period_end: date,
    snapshot: StatusSnapshot,
    generated_by: User | None,
    regenerate: bool = False,
) -> GeneratedReport:
    content, statistics = build_report_workbook(
        snapshot=snapshot, period_start=period_start, period_end=period_end
    )
    with transaction.atomic():
        last = (
            GeneratedReport.objects.select_for_update()
            .filter(report_type=report_type, report_date=report_date)
            .order_by("-version")
            .first()
        )
        if last and not regenerate:
            return last
        report = GeneratedReport(
            report_type=report_type,
            report_date=report_date,
            version=(last.version + 1) if last else 1,
            period_start=period_start,
            period_end=period_end,
            snapshot=snapshot,
            generated_by=generated_by,
            file_sha256=sha256(content).hexdigest(),
            statistics_range=statistics,
        )
        report.report_file.save(
            f"{report_type.lower()}-{report_date.isoformat()}-v{report.version}.xlsx",
            ContentFile(content),
            save=False,
        )
        report.save()
    return report
