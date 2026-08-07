import hashlib
import re
from collections.abc import Sequence
from datetime import date, datetime
from io import BytesIO
from typing import cast

from django.core.files.base import ContentFile
from django.db import IntegrityError, transaction
from django.db.models import QuerySet
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import generics, permissions, status
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import MultiPartParser
from rest_framework.permissions import BasePermission
from rest_framework.request import Request
from rest_framework.response import Response
from rest_framework.serializers import BaseSerializer
from rest_framework.views import APIView

from apps.accounts.models import User
from apps.audit.services import record_audit
from apps.core.permissions import CanOperate
from apps.core.views import ok

from .models import Article, ArticleImportJob, ArticleStatus, ImportStatus
from .serializers import ArticleImportJobSerializer, ArticleSerializer
from .services import normalize_title

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UPLOAD_ROWS = 1000
MAX_UPLOAD_COLUMNS = 10
MAX_CELL_CHARACTERS = 2048
ALLOWED_XLSX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
EXPECTED_HEADERS = {"原创文章标题", "原创发布日期", "原创文章链接", "原创发布平台", "作者或部门", "备注"}


class ArticleListCreateView(generics.ListCreateAPIView[Article]):
    serializer_class = ArticleSerializer

    def get_permissions(self) -> Sequence[BasePermission]:
        return [CanOperate()] if self.request.method == "POST" else [permissions.IsAuthenticated()]

    def get_queryset(self) -> QuerySet[Article]:
        queryset = Article.objects.select_related("created_by").all()
        if title := self.request.query_params.get("title"):
            queryset = queryset.filter(title__icontains=title)
        if status_value := self.request.query_params.get("status"):
            queryset = queryset.filter(status=status_value)
        if published_from := self.request.query_params.get("published_date_from"):
            queryset = queryset.filter(published_date__gte=published_from)
        if published_to := self.request.query_params.get("published_date_to"):
            queryset = queryset.filter(published_date__lte=published_to)
        return queryset

    def perform_create(self, serializer: BaseSerializer[Article]) -> None:
        article = cast(Article, serializer.save(created_by=cast(User, self.request.user)))
        record_audit(
            self.request,
            action_type="ARTICLE_CREATE",
            target_type="article",
            target_id=article.id,
            after_data=ArticleSerializer(article).data,
        )


class ArticleDetailView(generics.RetrieveUpdateAPIView[Article]):
    serializer_class = ArticleSerializer
    queryset = Article.objects.all()

    def get_permissions(self) -> Sequence[BasePermission]:
        return [CanOperate()] if self.request.method in {"PATCH", "PUT", "DELETE"} else [permissions.IsAuthenticated()]

    def perform_update(self, serializer: BaseSerializer[Article]) -> None:
        before = ArticleSerializer(self.get_object()).data
        article = cast(Article, serializer.save())
        record_audit(
            self.request,
            action_type="ARTICLE_UPDATE",
            target_type="article",
            target_id=article.id,
            before_data=before,
            after_data=ArticleSerializer(article).data,
        )

    def delete(self, request: Request, *args: object, **kwargs: object) -> Response:
        article = self.get_object()
        before = ArticleSerializer(article).data
        article.status = ArticleStatus.ARCHIVED
        article.save(update_fields=["status", "updated_at"])
        record_audit(
            request,
            action_type="ARTICLE_ARCHIVE",
            target_type="article",
            target_id=article.id,
            before_data=before,
            after_data=ArticleSerializer(article).data,
        )
        return Response(status=status.HTTP_204_NO_CONTENT)


class ArticleBulkStatusView(APIView):
    permission_classes = [CanOperate]
    serializer_class = ArticleSerializer

    def post(self, request: Request) -> Response:
        article_ids = request.data.get("article_ids")
        next_status = request.data.get("status")
        if not isinstance(article_ids, list) or not article_ids or len(article_ids) > 1000:
            raise ValidationError({"article_ids": "请提供 1 至 1000 个文章 ID。"})
        if next_status not in {ArticleStatus.ACTIVE, ArticleStatus.ARCHIVED}:
            raise ValidationError({"status": "批量操作仅支持归档或恢复。"})
        articles = list(Article.objects.filter(id__in=article_ids))
        changed_at = timezone.now()
        for article in articles:
            article.status = next_status
            article.updated_at = changed_at
        Article.objects.bulk_update(articles, ["status", "updated_at"])
        record_audit(
            request,
            action_type="ARTICLE_BULK_STATUS",
            target_type="article",
            after_data={"article_ids": [article.id for article in articles], "status": next_status},
        )
        return ok({"updated_count": len(articles), "status": next_status})


class ArticleBulkPasteView(APIView):
    permission_classes = [CanOperate]
    serializer_class = ArticleSerializer

    def post(self, request: Request) -> Response:
        items = request.data.get("articles")
        if not isinstance(items, list) or not items:
            raise ValidationError({"articles": "至少提供一篇文章。"})
        if len(items) > 100:
            raise ValidationError({"articles": "单次最多粘贴100篇文章。"})
        created: list[int] = []
        errors: list[dict[str, object]] = []
        with transaction.atomic():
            for index, item in enumerate(items, start=1):
                serializer = ArticleSerializer(data=item)
                if not serializer.is_valid():
                    errors.append({"row": index, "errors": serializer.errors})
                    continue
                article = serializer.save(created_by=request.user)
                created.append(article.id)
        record_audit(
            request,
            action_type="ARTICLE_BULK_CREATE",
            target_type="article",
            after_data={"created_count": len(created), "error_count": len(errors)},
        )
        return ok({"created_ids": created, "errors": errors}, status.HTTP_201_CREATED)


class ArticleImportPreviewView(APIView):
    permission_classes = [CanOperate]
    parser_classes = [MultiPartParser]
    serializer_class = ArticleImportJobSerializer

    def post(self, request: Request) -> Response:
        uploaded = request.FILES.get("file")
        if uploaded is None:
            raise ValidationError({"file": "请上传 .xlsx 文件。"})
        if (
            not uploaded.name.lower().endswith(".xlsx")
            or uploaded.content_type not in ALLOWED_XLSX_MIME_TYPES
            or uploaded.size > MAX_UPLOAD_BYTES
        ):
            raise ValidationError({"file": "只允许不超过10MB的 .xlsx 文件。"})
        raw = uploaded.read()
        if not raw.startswith(b"PK"):
            raise ValidationError({"file": "文件不是有效的 xlsx ZIP 容器。"})
        try:
            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        except Exception as error:
            raise ValidationError({"file": "无法安全读取该 xlsx 文件。"}) from error
        if len(workbook.worksheets) > 5:
            raise ValidationError({"file": "工作表数量超过限制。"})
        worksheet = workbook.active
        if worksheet.max_row > MAX_UPLOAD_ROWS + 1 or worksheet.max_column > MAX_UPLOAD_COLUMNS:
            raise ValidationError({"file": "文件行数或列数超过限制。"})
        headers = [str(cell.value or "").strip() for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
        if not {"原创文章标题", "原创发布日期"}.issubset(set(headers)):
            raise ValidationError({"file": "首行必须包含‘原创文章标题’和‘原创发布日期’。"})
        positions = {name: headers.index(name) for name in EXPECTED_HEADERS if name in headers}
        preview: list[dict[str, object]] = []
        for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            if any(cell is not None and len(str(cell)) > MAX_CELL_CHARACTERS for cell in row):
                preview.append({"row": row_number, "state": "FAILED", "reason": "单元格文本超过长度限制。", "data": {}})
                continue
            if any(isinstance(cell, str) and cell.startswith("=") for cell in row):
                preview.append({"row": row_number, "state": "FAILED", "reason": "不接受任何公式单元格。", "data": {}})
                continue
            title = str(row[positions["原创文章标题"]] or "").strip()
            raw_date = row[positions["原创发布日期"]]
            result: dict[str, object] = {"row": row_number, "title": title, "state": "FAILED", "reason": ""}
            if not title or not isinstance(raw_date, date):
                result["reason"] = "标题或发布日期无效。"
            elif Article.objects.filter(normalized_title=normalize_title(title), published_date=raw_date).exists():
                result["state"] = "DUPLICATE"
                result["reason"] = "同日相同标准化标题，需管理员确认。"
            else:
                result["state"] = "VALID"
            published_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            result["data"] = {
                "title": title,
                "published_date": published_date.isoformat() if isinstance(published_date, date) else "",
                "original_url": str(row[positions["原创文章链接"]] or "") if "原创文章链接" in positions else "",
                "source_platform": str(row[positions["原创发布平台"]] or "") if "原创发布平台" in positions else "",
                "author_department": str(row[positions["作者或部门"]] or "") if "作者或部门" in positions else "",
                "notes": str(row[positions["备注"]] or "") if "备注" in positions else "",
            }
            preview.append(result)
        job = ArticleImportJob(
            original_filename=re.split(r"[\\/]", uploaded.name)[-1][:255],
            sha256=hashlib.sha256(raw).hexdigest(),
            total_rows=len(preview),
            valid_rows=sum(item["state"] == "VALID" for item in preview),
            duplicate_rows=sum(item["state"] == "DUPLICATE" for item in preview),
            failed_rows=sum(item["state"] == "FAILED" for item in preview),
            preview_rows=preview,
            created_by=cast(User, request.user),
        )
        job.uploaded_file.save("upload.xlsx", ContentFile(raw), save=False)
        job.save()
        record_audit(
            request,
            action_type="ARTICLE_IMPORT_PREVIEW",
            target_type="article_import",
            target_id=job.id,
            after_data={"sha256": job.sha256, "total_rows": job.total_rows},
        )
        return ok(ArticleImportJobSerializer(job).data, status.HTTP_201_CREATED)


class ArticleImportConfirmView(APIView):
    permission_classes = [CanOperate]
    serializer_class = ArticleImportJobSerializer

    def post(self, request: Request, pk: int) -> Response:
        job = ArticleImportJob.objects.get(pk=pk)
        if job.status != ImportStatus.PREVIEW:
            raise ValidationError("该导入任务已处理。")
        confirm_duplicates = set(request.data.get("confirm_duplicate_rows", []))
        actor = cast(User, request.user)
        if confirm_duplicates and actor.role != "ADMIN":
            return Response(
                {"success": False, "error": {"code": "PERMISSION_DENIED", "message": "重复项需要管理员确认。"}},
                status=403,
            )
        imported: list[int] = []
        with transaction.atomic():
            for item in job.preview_rows:
                if item["state"] == "FAILED" or item["state"] == "DUPLICATE" and item["row"] not in confirm_duplicates:
                    continue
                data = item["data"]
                serializer = ArticleSerializer(data=data)
                if not serializer.is_valid():
                    continue
                try:
                    article = serializer.save(created_by=actor)
                except IntegrityError:
                    continue
                imported.append(article.id)
            job.status = ImportStatus.COMPLETED
            job.imported_article_ids = imported
            job.completed_at = timezone.now()
            job.save(update_fields=["status", "imported_article_ids", "completed_at"])
        record_audit(
            request,
            action_type="ARTICLE_IMPORT_CONFIRM",
            target_type="article_import",
            target_id=job.id,
            after_data={"imported_count": len(imported)},
        )
        return ok(ArticleImportJobSerializer(job).data)
