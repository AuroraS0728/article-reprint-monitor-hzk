from datetime import date
from io import BytesIO
from unittest.mock import Mock, patch

import pytest
from cryptography.fernet import Fernet
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import User
from apps.articles.models import Article
from apps.monitoring.models import PlatformDetectionStatus, SnapshotType, StatusSnapshot
from apps.platforms.models import Platform, PlatformStatus
from apps.reports.mail_services import encrypt_authorization_code, send_delivery
from apps.reports.models import DeliveryStatus, ReportEmailDelivery, ReportType, SMTPConfiguration
from apps.reports.services import build_report_workbook, create_report
from apps.reposts.models import RepostRecord
from tests.test_article_platform import client_for


def report_source(*, operator: User) -> tuple[StatusSnapshot, Article, Platform]:
    article = Article.objects.create(
        title="=external value",
        normalized_title="=external value",
        published_date=date(2026, 8, 3),
        created_by=operator,
    )
    platform = Platform.objects.create(code="TEST_REPORT_PLATFORM", name="测试平台", status=PlatformStatus.ENABLED)
    now = timezone.now()
    snapshot = StatusSnapshot.objects.create(
        snapshot_type=SnapshotType.WEEKLY,
        cutoff_at=now,
        matrix_data={
            f"{article.id}:{platform.id}": {
                "article_id": article.id,
                "platform_id": platform.id,
                "status": PlatformDetectionStatus.FOUND,
            }
        },
    )
    RepostRecord.objects.create(
        article=article,
        platform=platform,
        original_url="https://example.com/repost",
        normalized_url="https://example.com/repost",
        normalized_url_hash="a" * 64,
        final_url="https://example.com/repost",
        repost_title="=unsafe external title",
        first_discovered_at=now,
        last_checked_at=now,
        data_source="TEST_ONLY",
    )
    return snapshot, article, platform


@pytest.mark.django_db
def test_excel_report_uses_hyperlink_objects_and_sanitizes_formula_text(operator: User) -> None:
    snapshot, _, platform = report_source(operator=operator)
    content, statistics = build_report_workbook(
        snapshot=snapshot, period_start=date(2026, 8, 3), period_end=date(2026, 8, 9)
    )
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames == ["数据", platform.name, "新增转载链接", "总统计", "说明"]
    assert workbook["数据"]["B2"].value == "'=external value"
    assert workbook["数据"]["E2"].hyperlink is not None
    assert workbook[platform.name]["B2"].hyperlink.target == "https://example.com/repost"
    assert workbook["总统计"]["C2"].value == "√"
    assert statistics["new_repost_link_total"] == 1


@pytest.mark.django_db
def test_regenerated_report_creates_new_version(operator: User) -> None:
    snapshot, _, _ = report_source(operator=operator)
    first = create_report(
        report_type=ReportType.WEEKLY,
        report_date=date(2026, 8, 3),
        period_start=date(2026, 8, 3),
        period_end=date(2026, 8, 9),
        snapshot=snapshot,
        generated_by=operator,
    )
    same = create_report(
        report_type=ReportType.WEEKLY,
        report_date=date(2026, 8, 3),
        period_start=date(2026, 8, 3),
        period_end=date(2026, 8, 9),
        snapshot=snapshot,
        generated_by=operator,
    )
    regenerated = create_report(
        report_type=ReportType.WEEKLY,
        report_date=date(2026, 8, 3),
        period_start=date(2026, 8, 3),
        period_end=date(2026, 8, 9),
        snapshot=snapshot,
        generated_by=operator,
        regenerate=True,
    )
    assert same.id == first.id
    assert regenerated.version == 2


@pytest.mark.django_db
@override_settings(FIELD_ENCRYPTION_KEY=Fernet.generate_key().decode())
def test_email_delivery_uses_decrypted_secret_without_exposing_it(operator: User) -> None:
    snapshot, _, _ = report_source(operator=operator)
    report = create_report(
        report_type=ReportType.WEEKLY,
        report_date=date(2026, 8, 3),
        period_start=date(2026, 8, 3),
        period_end=date(2026, 8, 9),
        snapshot=snapshot,
        generated_by=operator,
    )
    SMTPConfiguration.objects.create(
        host="smtp.example.test",
        from_email="reports@example.test",
        encrypted_authorization_code=encrypt_authorization_code("not-a-real-secret"),
        recipients=["recipient@example.test"],
        enabled=True,
    )
    delivery = ReportEmailDelivery.objects.create(report=report, recipients=["recipient@example.test"])
    connection = Mock()
    message = Mock()
    with (
        patch("apps.reports.mail_services.get_connection", return_value=connection),
        patch("apps.reports.mail_services.EmailMessage", return_value=message),
    ):
        send_delivery(delivery)
    delivery.refresh_from_db()
    assert delivery.status == DeliveryStatus.SENT
    assert delivery.attempt_count == 1
    assert "not-a-real-secret" not in delivery.error_message


@pytest.mark.django_db
def test_report_download_and_regeneration_permissions(operator: User, administrator: User, viewer: User) -> None:
    snapshot, _, _ = report_source(operator=operator)
    report = create_report(
        report_type=ReportType.WEEKLY,
        report_date=date(2026, 8, 3),
        period_start=date(2026, 8, 3),
        period_end=date(2026, 8, 9),
        snapshot=snapshot,
        generated_by=operator,
    )
    assert client_for(viewer).get("/api/v1/reports").status_code == 200
    assert client_for(viewer).get(f"/api/v1/reports/{report.id}/download").status_code == 200
    assert client_for(viewer).post(f"/api/v1/reports/{report.id}/regenerate").status_code == 403
    response = client_for(administrator).post(f"/api/v1/reports/{report.id}/regenerate")
    assert response.status_code == 201
    assert response.json()["data"]["version"] == 2
