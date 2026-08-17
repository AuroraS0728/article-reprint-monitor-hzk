from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import timedelta
from hashlib import sha256
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from django.core.management import call_command
from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.articles.models import Article, ArticleIngestMethod, ArticleMonitoringStatus
from apps.audit.models import OperationLog
from apps.reposts.export_services import (
    EXCEL_DATETIME_FORMAT,
    GlobalExportFilters,
    build_global_repost_workbook,
    filtered_articles,
    sanitize_excel_sheet_name,
)
from apps.reposts.models import ContentRelation, RepostRecord
from apps.reposts.query_services import (
    RepostQueryFilters,
    filtered_repost_articles,
    repost_trend,
    result_summary,
)
from apps.sources.models import (
    AutomaticRepostSite,
    OwnedChannel,
    SearchCandidateDisposition,
    SearchProviderConfiguration,
    SearchRun,
    SearchRunCandidate,
    SearchRunStatus,
    Source,
)
from apps.sources.reading_export_services import build_owned_reading_workbook
from apps.sources.services import (
    _next_search_time,
    canonicalize_http_url,
    compare_titles,
    ingest_source_article,
    mark_repost_availability,
    purge_expired_source_articles,
    search_article,
)
from apps.sources.tasks import search_article_reposts
from apps.sources.token_services import create_source_token
from search_providers.exceptions import SearchProviderRateLimited
from search_providers.types import SearchCandidate


@dataclass
class StaticProvider:
    candidates: list[SearchCandidate]
    code: str = "test-provider"

    def search(self, query: str, **kwargs: object) -> list[SearchCandidate]:
        return [
            (
                candidate
                if candidate.published_at is not None
                else replace(candidate, published_at=timezone.now() + timedelta(minutes=1))
            )
            for candidate in self.candidates
        ]


@dataclass
class RawStaticProvider:
    """Returns provider data unchanged for retention-gate tests."""

    candidates: list[SearchCandidate]
    code: str = "raw-test-provider"

    def search(self, query: str, **kwargs: object) -> list[SearchCandidate]:
        return list(self.candidates)


class ErrorProvider:
    code = "test-provider"

    def search(self, query: str, **kwargs: object) -> list[SearchCandidate]:
        raise SearchProviderRateLimited("rate limited")


class NamedStaticProvider(StaticProvider):
    def __init__(self, code: str, candidates: list[SearchCandidate]) -> None:
        super().__init__(candidates=candidates, code=code)


class PerStageProvider:
    code = "per-stage"

    def __init__(self, responses: dict[str, list[SearchCandidate] | Exception]) -> None:
        self.responses = responses
        self.calls: list[str] = []

    def search(self, query: str, **kwargs: object) -> list[SearchCandidate]:
        self.calls.append(query)
        response = self.responses[query]
        if isinstance(response, Exception):
            raise response
        return [
            (
                candidate
                if candidate.published_at is not None
                else replace(candidate, published_at=timezone.now() + timedelta(minutes=1))
            )
            for candidate in response
        ]


class InspectingProvider:
    """Confirms that the first provider response is saved before the next query."""

    code = "inspecting-provider"

    def __init__(self) -> None:
        self.calls = 0

    def search(self, query: str, **kwargs: object) -> list[SearchCandidate]:
        self.calls += 1
        if self.calls == 1:
            return [
                SearchCandidate(
                    title="第一批候选",
                    url="https://example.com/first",
                    site_name="示例站",
                    published_at=timezone.now() + timedelta(minutes=1),
                )
            ]

        run = SearchRun.objects.latest("id")
        candidate = SearchRunCandidate.objects.get(search_run=run)
        assert run.status == SearchRunStatus.RUNNING
        assert run.candidate_count == 1
        assert candidate.disposition == SearchCandidateDisposition.PENDING
        return [
            SearchCandidate(
                title="第二批候选",
                url="https://example.com/second",
                site_name="示例站",
                published_at=timezone.now() + timedelta(minutes=1),
            )
        ]


@pytest.fixture
def source(db: object) -> Source:
    source, _ = Source.objects.update_or_create(
        code="WEEKLYONSTOCK",
        defaults={
            "name": "证券市场周刊",
            "base_url": "https://www.weeklyonstock.com/",
            "is_active": True,
        },
    )
    return source


def create_source_article(*, source: Source, operator: User, days_old: int = 0) -> Article:
    published_at = timezone.now() - timedelta(days=days_old)
    return ingest_source_article(
        source=source,
        title="这一主线再掀涨停潮！低位方向正在成为新主角？",
        author="木木",
        published_at=published_at,
        original_url="https://www.weeklyonstock.com/article/123?utm_source=feed",
        created_by=operator,
    ).article


@pytest.mark.django_db
def test_source_token_is_hashed_and_ingest_api_is_idempotent(source: Source, operator: User) -> None:
    created_token = create_source_token(source=source, name="browser", created_by=operator)
    assert created_token.record.token_hash != created_token.plaintext
    assert created_token.plaintext not in created_token.record.token_prefix
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {created_token.plaintext}")
    payload = {
        "title": "原创标题",
        "author": "作者",
        "published_at": timezone.now().isoformat(),
        "original_url": "https://www.weeklyonstock.com/article/ingest?utm_source=browser",
    }
    with patch("apps.sources.tasks.search_article_reposts.delay"):
        first = client.post("/api/v1/source-ingest/articles", payload, format="json")
        second = client.post("/api/v1/source-ingest/articles", payload, format="json")
        payload["title"] = "修改后的原创标题"
        third = client.post("/api/v1/source-ingest/articles", payload, format="json")
    assert first.status_code == 201
    assert second.status_code == 200
    assert second.json()["data"]["updated"] is False
    assert third.status_code == 200
    assert third.json()["data"]["updated"] is True
    assert Article.objects.filter(source=source).count() == 1
    article = Article.objects.get(source=source)
    assert article.title == "修改后的原创标题"
    assert article.ingest_method == ArticleIngestMethod.SOURCE_API


@pytest.mark.django_db
def test_source_ingest_accepts_optional_channel_fields_without_breaking_existing_payload(
    source: Source, operator: User
) -> None:
    token = create_source_token(source=source, name="channel-test", created_by=operator)
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION=f"Bearer {token.plaintext}")
    payload = {
        "title": "带栏目字段的原创文章",
        "author": "作者",
        "published_at": timezone.now().isoformat(),
        "original_url": "https://www.weeklyonstock.com/article/channel-test",
        "channel_code": "RIGHTS",
        "channel_name": "证券维权",
        "section_code": "COMPANY",
        "section_name": "涉事公司报道",
    }
    with patch("apps.sources.tasks.search_article_reposts.delay"):
        response = client.post("/api/v1/source-ingest/articles", payload, format="json")
    assert response.status_code == 201
    article = Article.objects.get(pk=response.json()["data"]["article_id"])
    assert (article.channel_code, article.section_code) == ("RIGHTS", "COMPANY")


@pytest.mark.django_db
def test_source_ingest_rejects_invalid_expired_inactive_tokens_and_source(source: Source, operator: User) -> None:
    client = APIClient()
    client.credentials(HTTP_AUTHORIZATION="Bearer invalid")
    assert client.post("/api/v1/source-ingest/articles", {}, format="json").status_code == 401

    payload = {
        "title": "原创标题",
        "author": "作者",
        "published_at": timezone.now().isoformat(),
        "original_url": "https://www.weeklyonstock.com/a",
    }
    for token_changes, source_active in [
        ({"expires_at": timezone.now() - timedelta(seconds=1)}, True),
        ({"is_active": False}, True),
        ({}, False),
    ]:
        source.is_active = source_active
        source.save(update_fields=["is_active"])
        created = create_source_token(source=source, name="test", created_by=operator)
        for field, value in token_changes.items():
            setattr(created.record, field, value)
        created.record.save()
        client.credentials(HTTP_AUTHORIZATION=f"Bearer {created.plaintext}")
        assert client.post("/api/v1/source-ingest/articles", payload, format="json").status_code == 401


@pytest.mark.django_db
def test_expired_first_ingest_is_completed_without_restarting_monitoring(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator, days_old=8)
    assert article.monitoring_status == ArticleMonitoringStatus.COMPLETED
    assert article.next_search_at is None
    original_until = article.monitor_until
    result = ingest_source_article(
        source=source,
        title=f"{article.title} 修改",
        author="木木",
        published_at=article.published_at,
        original_url=article.original_url,
        created_by=operator,
    )
    assert result.article.monitor_until == original_until
    assert result.article.monitoring_status == ArticleMonitoringStatus.COMPLETED


@pytest.mark.django_db
def test_manual_global_search_allows_completed_article_without_restarting_lifecycle(
    source: Source, operator: User
) -> None:
    article = create_source_article(source=source, operator=operator, days_old=8)
    original_next_search_at = article.next_search_at
    with (
        patch("apps.sources.tasks.search_article", return_value=SimpleNamespace(id=77, status="SUCCESS")) as search,
        patch("apps.sources.tasks.redis.Redis.from_url") as redis_from_url,
    ):
        lock = redis_from_url.return_value.lock.return_value
        lock.acquire.return_value = True
        assert search_article_reposts.run(article.id, manual=True) == 77
    article.refresh_from_db()
    assert search.called
    assert article.monitoring_status == ArticleMonitoringStatus.COMPLETED
    assert article.next_search_at == original_next_search_at


@pytest.mark.django_db
def test_manual_global_search_api_queues_unarchived_articles_and_rejects_archived(
    source: Source, operator: User
) -> None:
    article = create_source_article(source=source, operator=operator)
    client = APIClient()
    client.force_authenticate(operator)
    with patch("apps.sources.views.search_article_reposts.delay") as delay:
        response = client.post("/api/v1/global-search-runs", {"article_ids": [article.id]}, format="json")
    assert response.status_code == 202
    assert response.json()["data"]["queued_count"] == 1
    delay.assert_called_once_with(article.id, manual=True)
    article.status = "ARCHIVED"
    article.save(update_fields=["status"])
    response = client.post("/api/v1/global-search-runs", {"article_ids": [article.id]}, format="json")
    assert response.status_code == 400


@pytest.mark.django_db
def test_search_run_candidates_are_available_to_authenticated_readers(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    run = search_article(
        article,
        provider=StaticProvider(
            [SearchCandidate(title=article.title, url="https://finance.example.com/a", site_name="示例财经")]
        ),
    )
    client = APIClient()
    client.force_authenticate(operator)
    response = client.get(f"/api/v1/global-search-runs/{run.id}/candidates")
    assert response.status_code == 200
    candidate = response.json()["data"][0]
    assert candidate["title"] == article.title
    assert candidate["site_name"] == "示例财经"
    assert candidate["canonical_url"] == "https://finance.example.com/a"


@pytest.mark.django_db
def test_search_retains_high_similarity_candidates_without_publication_time(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    later = article.published_at + timedelta(minutes=1)
    run = search_article(
        article,
        provider=RawStaticProvider(
            [
                SearchCandidate(title=article.title, url="https://news.example.com/late", published_at=later),
                SearchCandidate(
                    title=article.title,
                    url="https://news.example.com/same-time",
                    published_at=article.published_at,
                ),
                SearchCandidate(title=article.title, url="https://news.example.com/no-time"),
                SearchCandidate(
                    title="与原创标题没有关联的搜索结果",
                    url="https://news.example.com/unrelated",
                    published_at=later,
                ),
            ]
        ),
    )

    retained = SearchRunCandidate.objects.filter(search_run=run)
    assert list(retained.values_list("canonical_url", flat=True)) == [
        "https://news.example.com/late",
        "https://news.example.com/same-time",
        "https://news.example.com/no-time",
    ]
    assert run.exact_candidate_count == 3
    assert run.broad_candidate_count == 3
    assert run.merged_candidate_count == 3

    client = APIClient()
    client.force_authenticate(operator)
    response = client.get(f"/api/v1/articles/{article.id}/discovered-publications")

    assert response.status_code == 200
    assert {row["canonical_url"] for row in response.json()["data"]["candidates"]} == {
        "https://news.example.com/late",
        "https://news.example.com/same-time",
        "https://news.example.com/no-time",
    }


@pytest.mark.django_db
def test_article_candidate_detail_shows_one_latest_row_per_canonical_url(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    candidate = SearchCandidate(
        title=article.title,
        url="https://news.example.com/only-once?utm_source=first",
        published_at=article.published_at + timedelta(minutes=1),
    )
    first_run = search_article(article, provider=StaticProvider([candidate]))
    second_run = search_article(article, provider=StaticProvider([candidate]))
    client = APIClient()
    client.force_authenticate(operator)

    response = client.get(f"/api/v1/articles/{article.id}/discovered-publications")

    assert response.status_code == 200
    rows = response.json()["data"]["candidates"]
    assert [row["canonical_url"] for row in rows] == ["https://news.example.com/only-once"]
    assert rows[0]["search_run_id"] == second_run.id
    assert SearchRunCandidate.objects.filter(search_run__in=[first_run, second_run]).count() == 2


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=101)
def test_operator_can_review_candidate_as_external_repost_and_audit_is_recorded(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    run = search_article(
        article,
        provider=StaticProvider(
            [
                SearchCandidate(
                    title=article.title,
                    url="https://manual-review.example.com/a",
                    site_name="人工复核站",
                )
            ]
        ),
    )
    candidate = SearchRunCandidate.objects.get(search_run=run)
    assert not RepostRecord.objects.filter(article=article).exists()

    client = APIClient()
    client.force_authenticate(operator)
    response = client.post(
        f"/api/v1/search-candidates/{candidate.id}/review",
        {"action": "CONFIRM_REPOST", "reason": "人工确认是外部转载"},
        format="json",
    )

    assert response.status_code == 200
    candidate.refresh_from_db()
    record = RepostRecord.objects.get(article=article, canonical_url_hash=candidate.canonical_url_hash)
    assert candidate.content_relation == ContentRelation.REPOST
    assert candidate.disposition == SearchCandidateDisposition.MATCHED
    assert record.content_relation == ContentRelation.REPOST
    assert run.__class__.objects.get(pk=run.id).repost_count == 1
    assert OperationLog.objects.filter(action_type="SEARCH_CANDIDATE_REVIEWED", target_id=str(candidate.id)).exists()


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=101)
def test_operator_can_send_candidate_to_reading_module_without_repost_export(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    channel = OwnedChannel.objects.create(
        code="OWNED_REVIEW",
        name="人工复核自有渠道",
        channel_type="OFFICIAL_WEBSITE",
        match_rules={"domains": ["owned-review.example.com"]},
    )
    run = search_article(
        article,
        provider=StaticProvider(
            [SearchCandidate(title=article.title, url="https://owned-review.example.com/a", site_name="自有渠道")]
        ),
    )
    candidate = SearchRunCandidate.objects.get(search_run=run)
    client = APIClient()
    client.force_authenticate(operator)

    response = client.post(
        f"/api/v1/search-candidates/{candidate.id}/review",
        {"action": "CONFIRM_OWNED", "reason": "属于自有分发", "owned_channel_id": channel.id},
        format="json",
    )

    assert response.status_code == 200
    candidate.refresh_from_db()
    record = RepostRecord.objects.get(article=article, canonical_url_hash=candidate.canonical_url_hash)
    assert candidate.content_relation == ContentRelation.OWNED
    assert candidate.owned_channel_id == channel.id
    assert record.content_relation == ContentRelation.OWNED
    assert record.owned_channel_id == channel.id
    as_of = timezone.now()
    assert (
        result_summary(filtered_repost_articles(RepostQueryFilters(), as_of=as_of), as_of=as_of)["repost_url_count"]
        == 0
    )
    reading = client.get("/api/v1/reading-monitor/owned-publications")
    assert reading.status_code == 200
    assert reading.json()["data"]["publications"][0]["url"] == "https://owned-review.example.com/a"

    content, _ = build_owned_reading_workbook()
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames[0] == "数据"
    assert channel.name in workbook.sheetnames
    assert workbook.sheetnames[-2:] == ["总统计", "说明"]
    assert workbook["数据"]["D2"].value == "—"
    assert workbook[channel.name]["D2"].hyperlink.target == "https://owned-review.example.com/a"
    assert workbook["总统计"]["D2"].value == 0
    assert "外部转载" in workbook["说明"]["B2"].value


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=101)
def test_operator_can_switch_between_manual_repost_and_owned_reviews(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    channel = OwnedChannel.objects.create(
        code="CHANGE_REVIEW_TO_OWNED",
        name="修改复核自有渠道",
        channel_type="PLATFORM_ACCOUNT",
        match_rules={"platform_domains": ["change-review.example.com"]},
    )
    run = search_article(
        article,
        provider=StaticProvider([SearchCandidate(title=article.title, url="https://change-review.example.com/a")]),
    )
    candidate = SearchRunCandidate.objects.get(search_run=run)
    client = APIClient()
    client.force_authenticate(operator)

    confirmed = client.post(
        f"/api/v1/search-candidates/{candidate.id}/review",
        {"action": "CONFIRM_REPOST", "reason": "人工确认外部转载"},
        format="json",
    )
    changed = client.post(
        f"/api/v1/search-candidates/{candidate.id}/review",
        {"action": "CONFIRM_OWNED", "reason": "复核后改为自有渠道", "owned_channel_id": channel.id},
        format="json",
    )

    assert confirmed.status_code == 200
    assert changed.status_code == 200
    candidate.refresh_from_db()
    record = RepostRecord.objects.get(article=article, canonical_url_hash=candidate.canonical_url_hash)
    assert candidate.reason_code == "MANUAL_OWNED"
    assert candidate.content_relation == ContentRelation.OWNED
    assert candidate.owned_channel_id == channel.id
    assert record.content_relation == ContentRelation.OWNED
    assert record.owned_channel_id == channel.id

    changed_back = client.post(
        f"/api/v1/search-candidates/{candidate.id}/review",
        {"action": "CONFIRM_REPOST", "reason": "再次复核后改回外部转载"},
        format="json",
    )

    assert changed_back.status_code == 200
    candidate.refresh_from_db()
    record.refresh_from_db()
    assert candidate.reason_code == "MANUAL_REPOST"
    assert candidate.content_relation == ContentRelation.REPOST
    assert candidate.owned_channel_id is None
    assert record.content_relation == ContentRelation.REPOST
    assert record.owned_channel_id is None
    detail = client.get(f"/api/v1/articles/{article.id}/discovered-publications")
    assert detail.status_code == 200
    row = detail.json()["data"]["candidates"][0]
    assert row["manual_review_action"] == "CONFIRM_REPOST"
    assert row["owned_channel_id"] is None
    audits = OperationLog.objects.filter(action_type="SEARCH_CANDIDATE_REVIEWED", target_id=str(candidate.id))
    assert audits.count() == 3
    assert audits.order_by("id").last().after_data["content_relation"] == ContentRelation.REPOST


@pytest.mark.django_db
def test_reading_module_exposes_manual_other_owned_channel(operator: User) -> None:
    client = APIClient()
    client.force_authenticate(operator)

    response = client.get("/api/v1/reading-monitor/owned-publications")

    assert response.status_code == 200
    channels = response.json()["data"]["channels"]
    assert {channel["code"] for channel in channels} >= {"OTHER"}
    assert next(channel for channel in channels if channel["code"] == "OTHER")["name"] == "其他"


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=101)
def test_candidate_review_requires_operator_or_administrator(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    run = search_article(
        article,
        provider=StaticProvider([SearchCandidate(title=article.title, url="https://review-auth.example.com/a")]),
    )
    viewer = User.objects.create_user(username="candidate-viewer", password="A-valid-password-123", role="VIEWER")
    client = APIClient()
    client.force_authenticate(viewer)
    response = client.post(
        f"/api/v1/search-candidates/{SearchRunCandidate.objects.get(search_run=run).id}/review",
        {"action": "EXCLUDE", "reason": "只读用户不可复核"},
        format="json",
    )
    assert response.status_code == 403


@pytest.mark.django_db
@override_settings(SEARCH_CANDIDATE_MIN_SIMILARITY=0)
def test_search_candidates_are_persisted_while_run_is_still_running(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    run = search_article(article, provider=InspectingProvider())

    assert run.status == SearchRunStatus.SUCCESS
    assert run.candidate_count == 2
    assert SearchRunCandidate.objects.filter(search_run=run).count() == 2


def test_url_canonicalization_and_similarity_threshold() -> None:
    assert canonicalize_http_url("HTTPS://Example.COM:443/a?utm_source=x&a=1#fragment") == "https://example.com/a?a=1"
    with pytest.raises(ValueError):
        canonicalize_http_url("http://127.0.0.1/private")
    exact = compare_titles(
        "标题：测试！",
        SearchCandidate(title="标题测试 - 某网站", url="https://example.com/1", site_name="某网站"),
    )
    assert exact.matched is True
    assert exact.similarity_score == 100
    with override_settings(SEARCH_SHORT_TITLE_LENGTH=1, SEARCH_SIMILARITY_THRESHOLD=90):
        with patch("apps.sources.services.fuzz.ratio", return_value=89.99):
            assert (
                compare_titles("足够长的测试标题", SearchCandidate(title="另一标题", url="https://a.com")).matched
                is False
            )
        with patch("apps.sources.services.fuzz.ratio", return_value=90.0):
            assert (
                compare_titles("足够长的测试标题", SearchCandidate(title="另一标题", url="https://a.com")).matched
                is True
            )


@pytest.mark.django_db
def test_global_search_deduplicates_urls_excludes_source_and_preserves_history(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    title = article.title
    candidates = [
        SearchCandidate(title=title, url="https://news.example.com/1?a=1&utm_source=x", site_name="示例财经"),
        SearchCandidate(title=title, url="https://news.example.com/1?utm_source=y&a=1", site_name="示例财经"),
        SearchCandidate(title=title, url="https://news.example.com/2", site_name="示例财经"),
        SearchCandidate(title=title, url="https://static.weeklyonstock.com/copy", site_name="原站"),
    ]
    first_run = search_article(article, provider=StaticProvider(candidates))
    assert first_run.status == SearchRunStatus.SUCCESS
    assert SearchRunCandidate.objects.filter(search_run=first_run).count() == 3
    owned_candidate = SearchRunCandidate.objects.get(search_run=first_run, canonical_url__contains="weeklyonstock")
    assert owned_candidate.content_relation == ContentRelation.OWNED
    assert RepostRecord.objects.filter(article=article).count() == 3
    assert RepostRecord.objects.filter(article=article, content_relation=ContentRelation.REPOST).count() == 2
    first_record = RepostRecord.objects.filter(article=article).order_by("canonical_url").first()
    assert first_record is not None
    first_found = first_record.first_found_at

    search_article(article, provider=StaticProvider([]))
    assert RepostRecord.objects.filter(article=article).count() == 3
    mark_repost_availability(record=first_record, status="REMOVED")
    first_record.refresh_from_db()
    assert first_record.first_found_at == first_found
    assert first_record.availability_status == "REMOVED"
    assert (
        RepostRecord.objects.filter(
            article=article,
            is_valid=True,
            content_relation=ContentRelation.REPOST,
        ).count()
        == 2
    )


@pytest.mark.django_db
def test_two_stage_search_runs_exact_then_broad_deduplicates_and_retains_stage_counts(
    source: Source, operator: User
) -> None:
    article = create_source_article(source=source, operator=operator)
    exact_query = f'"{article.title}"'
    broad_query = article.title
    shared = SearchCandidate(title=article.title, url="https://external.example.com/shared")
    broad_only = SearchCandidate(title=article.title, url="https://external.example.com/broad")
    provider = PerStageProvider({exact_query: [shared], broad_query: [shared, broad_only]})

    run = search_article(article, provider=provider)

    assert provider.calls == [exact_query, broad_query]
    assert run.exact_candidate_count == 1
    assert run.broad_candidate_count == 2
    assert run.merged_candidate_count == 2
    assert run.repost_count == 2
    shared_candidate = SearchRunCandidate.objects.get(
        search_run=run, canonical_url="https://external.example.com/shared"
    )
    assert shared_candidate.search_phases == ["EXACT", "BROAD"]


@pytest.mark.django_db
def test_broad_stage_can_succeed_after_exact_stage_failure(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    provider = PerStageProvider(
        {
            f'"{article.title}"': SearchProviderRateLimited("exact limited"),
            article.title: [SearchCandidate(title=article.title, url="https://external.example.com/fallback")],
        }
    )

    run = search_article(article, provider=provider)

    assert run.status == SearchRunStatus.SUCCESS
    assert run.error_code == "PARTIAL_STAGE_FAILURE"
    assert run.repost_count == 1


@pytest.mark.django_db
def test_enabled_search_providers_run_in_priority_order_and_fall_back_after_failure(
    source: Source, operator: User
) -> None:
    article = create_source_article(source=source, operator=operator)
    first = SearchProviderConfiguration.objects.create(code="tencent_wsa", name="First", enabled=True, priority=10)
    second = SearchProviderConfiguration.objects.create(code="brave", name="Second", enabled=True, priority=20)
    SearchProviderConfiguration.objects.filter(code="bing_html").update(enabled=False)
    fallback = NamedStaticProvider(
        "brave",
        [SearchCandidate(title=article.title, url="https://fallback.example.com/repost", site_name="Fallback")],
    )
    with patch(
        "apps.sources.services.search_provider_for_code",
        side_effect=[ErrorProvider(), fallback],
    ):
        run = search_article(article)

    first.refresh_from_db()
    second.refresh_from_db()
    assert run.status == SearchRunStatus.SUCCESS
    assert run.provider == "tencent_wsa,brave"
    candidate = SearchRunCandidate.objects.get(search_run=run, canonical_url="https://fallback.example.com/repost")
    assert candidate.provider_codes == ["brave"]
    assert first.consecutive_failures == 1
    assert first.last_failure_code == "SEARCH_PROVIDER_RATE_LIMITED"
    assert second.consecutive_failures == 0
    assert second.last_success_at is not None


@pytest.mark.django_db
def test_candidate_keeps_all_provider_codes_when_the_same_url_is_returned(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    first = SearchProviderConfiguration.objects.create(code="tencent_wsa", name="First", enabled=True, priority=10)
    second = SearchProviderConfiguration.objects.create(code="brave", name="Second", enabled=True, priority=20)
    SearchProviderConfiguration.objects.filter(code="bing_html").update(enabled=False)
    shared_url = "https://same.example.com/repost"
    with patch(
        "apps.sources.services.search_provider_for_code",
        side_effect=[
            NamedStaticProvider("tencent_wsa", [SearchCandidate(title=article.title, url=shared_url)]),
            NamedStaticProvider("brave", [SearchCandidate(title=article.title, url=shared_url)]),
        ],
    ):
        run = search_article(article)

    candidate = SearchRunCandidate.objects.get(search_run=run, canonical_url=shared_url)
    assert candidate.provider_codes == ["tencent_wsa", "brave"]
    first.refresh_from_db()
    second.refresh_from_db()
    assert first.last_success_at is not None
    assert second.last_success_at is not None


@pytest.mark.django_db
def test_search_provider_configuration_api_requires_admin_and_configured_secret(source: Source, operator: User) -> None:
    client = APIClient()
    client.force_authenticate(operator)
    forbidden = client.get("/api/v1/search-providers")
    assert forbidden.status_code == 403

    admin = User.objects.create_user(username="search-admin", password="A-valid-password-123", role="ADMIN")
    client.force_authenticate(admin)
    with override_settings(BRAVE_SEARCH_API_KEY=""):
        rejected = client.post(
            "/api/v1/search-providers",
            {"code": "brave", "name": "Brave", "enabled": True, "priority": 10},
            format="json",
        )
    assert rejected.status_code == 400
    with override_settings(TENCENTCLOUD_WSA_APIKEY="test-key"):
        created = client.post(
            "/api/v1/search-providers",
            {"code": "tencent_wsa", "name": "腾讯云联网搜索", "enabled": True, "priority": 10},
            format="json",
        )
    assert created.status_code == 201
    assert OperationLog.objects.filter(action_type="SEARCH_PROVIDER_CONFIGURATION_CREATE").exists()


@pytest.mark.django_db
def test_owned_channel_requires_configured_account_evidence_and_external_is_repost(
    source: Source, operator: User
) -> None:
    article = create_source_article(source=source, operator=operator)
    OwnedChannel.objects.create(
        code="TEST_PLATFORM",
        name="测试自有账号",
        channel_type="PLATFORM_ACCOUNT",
        match_rules={"platform_domains": ["platform.example.com"], "account_names": ["证券市场周刊"]},
    )
    provider = StaticProvider(
        [
            SearchCandidate(
                title=article.title,
                url="https://platform.example.com/owned",
                raw_data={"account_name": "证券市场周刊"},
            ),
            SearchCandidate(title=article.title, url="https://platform.example.com/uncertain"),
            SearchCandidate(title=article.title, url="https://external.example.com/repost"),
        ]
    )

    run = search_article(article, provider=provider)

    assert run.owned_count == 1
    assert run.review_required_count == 1
    assert run.repost_count == 1
    assert RepostRecord.objects.filter(article=article, content_relation=ContentRelation.OWNED).count() == 1
    assert RepostRecord.objects.filter(article=article, content_relation=ContentRelation.REPOST).count() == 1


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=90, SEARCH_CANDIDATE_MIN_SIMILARITY=80)
def test_configured_media_site_auto_confirms_candidate_threshold_and_keeps_confirmed_owned_priority(
    source: Source, operator: User
) -> None:
    article = create_source_article(source=source, operator=operator)
    AutomaticRepostSite.objects.create(
        code="AUTO_MEDIA",
        name="自动转载媒体",
        domains=["auto-media.example.com"],
    )
    OwnedChannel.objects.create(
        code="AUTO_MEDIA_OWNED",
        name="自动转载媒体自有账号",
        channel_type="PLATFORM_ACCOUNT",
        match_rules={"platform_domains": ["auto-media.example.com"], "account_names": ["证券市场周刊"]},
    )
    external = SearchCandidate(
        title="这一主线再掀涨停潮低位方向正在成为新主角报道",
        url="https://auto-media.example.com/repost",
        site_name="自动转载媒体",
    )
    owned = SearchCandidate(
        title="这一主线再掀涨停潮低位方向正在成为新主角报道",
        url="https://auto-media.example.com/owned",
        site_name="自动转载媒体",
        raw_data={"account_name": "证券市场周刊"},
    )

    with patch("apps.sources.services.fuzz.ratio", return_value=85.0):
        run = search_article(article, provider=RawStaticProvider([external, owned]))

    repost = RepostRecord.objects.get(article=article, canonical_url__contains="/repost")
    assert repost.content_relation == ContentRelation.REPOST
    assert repost.classification_reason == "AUTO_REPOST_SITE_DOMAIN:AUTO_MEDIA"
    # A positively identified first-party account enters reading measurement,
    # and is never counted as an external repost.
    owned_record = RepostRecord.objects.get(article=article, canonical_url__contains="/owned")
    assert owned_record.content_relation == ContentRelation.OWNED
    assert owned_record.owned_channel_id is not None
    assert run.repost_count == 1
    assert run.owned_count == 1
    assert run.matched_count == 2


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=90, SEARCH_CANDIDATE_MIN_SIMILARITY=80)
def test_auto_repost_reclassification_promotes_retained_historical_candidates(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    candidate = SearchCandidate(
        title="这一主线再掀涨停潮低位方向正在成为新主角报道",
        url="https://historical-auto-media.example.com/repost",
    )
    with patch("apps.sources.services.fuzz.ratio", return_value=85.0):
        run = search_article(article, provider=RawStaticProvider([candidate]))
    assert not RepostRecord.objects.filter(article=article).exists()

    AutomaticRepostSite.objects.create(
        code="HISTORICAL_AUTO_MEDIA",
        name="历史自动转载媒体",
        domains=["historical-auto-media.example.com"],
    )
    output = StringIO()
    with patch("apps.sources.services.fuzz.ratio", return_value=85.0):
        call_command("reclassify_automatic_repost_sites", stdout=output)

    saved = RepostRecord.objects.get(article=article)
    assert saved.content_relation == ContentRelation.REPOST
    assert saved.classification_reason == "AUTO_REPOST_SITE_DOMAIN:HISTORICAL_AUTO_MEDIA"
    assert SearchRunCandidate.objects.get(search_run=run).content_relation == ContentRelation.REPOST
    assert "APPLIED: changed=1" in output.getvalue()


@pytest.mark.django_db
def test_reclassification_never_confirms_or_keeps_automatic_candidate_below_retention_threshold(
    source: Source, operator: User
) -> None:
    article = create_source_article(source=source, operator=operator)
    AutomaticRepostSite.objects.create(
        code="LOW_SCORE_AUTO_MEDIA",
        name="低分自动转载媒体",
        domains=["low-score-auto-media.example.com"],
    )
    candidate = SearchCandidate(
        title="与原标题只有少量共同词的无关新闻",
        url="https://low-score-auto-media.example.com/unrelated",
    )
    # Reproduce a legacy incorrect record made while the retention gate was too low.
    with override_settings(SEARCH_CANDIDATE_MIN_SIMILARITY=0):
        with patch("apps.sources.services.fuzz.ratio", return_value=75.0):
            run = search_article(article, provider=RawStaticProvider([candidate]))
    assert RepostRecord.objects.filter(article=article).exists()
    assert SearchRunCandidate.objects.filter(search_run=run).exists()

    output = StringIO()
    with override_settings(SEARCH_CANDIDATE_MIN_SIMILARITY=80):
        with patch("apps.sources.services.fuzz.ratio", return_value=75.0):
            call_command("reclassify_automatic_repost_sites", "--prune-below-threshold", stdout=output)

    assert not RepostRecord.objects.filter(article=article).exists()
    assert not SearchRunCandidate.objects.filter(search_run=run).exists()
    assert "below_threshold_pruned=1" in output.getvalue()


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=90, SEARCH_CANDIDATE_MIN_SIMILARITY=80)
def test_caifuhao_subdomain_is_auto_classified_as_owned_reading_channel(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    candidate = SearchCandidate(
        title="这一主线再掀涨停潮低位方向正在成为新主角报道",
        url="https://caifuhao.eastmoney.com/news/example",
        site_name="财富号",
    )
    owned_channel = OwnedChannel.objects.get(code="EASTMONEY")
    assert owned_channel.name == "财富号"
    assert owned_channel.match_rules == {"subdomains": ["caifuhao.eastmoney.com"]}

    with patch("apps.sources.services.fuzz.ratio", return_value=85.0):
        run = search_article(article, provider=RawStaticProvider([candidate]))

    saved = RepostRecord.objects.get(article=article)
    assert saved.content_relation == ContentRelation.OWNED
    assert saved.owned_channel is not None
    assert saved.owned_channel.code == "EASTMONEY"
    assert saved.classification_reason == "OWNED_CHANNEL_DOMAIN_MATCH"
    assert run.owned_count == 1
    assert run.repost_count == 0


@pytest.mark.django_db
@override_settings(SEARCH_SIMILARITY_THRESHOLD=90, SEARCH_CANDIDATE_MIN_SIMILARITY=80)
def test_reclassification_moves_historical_caifuhao_candidate_to_owned_reading(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    candidate = SearchCandidate(
        title="这一主线再掀涨停潮低位方向正在成为新主角报道",
        url="https://caifuhao.eastmoney.com/news/historical",
        site_name="财富号",
    )
    owned_channel = OwnedChannel.objects.get(code="EASTMONEY")
    owned_channel.is_active = False
    owned_channel.save(update_fields=["is_active"])
    with patch("apps.sources.services.fuzz.ratio", return_value=85.0):
        run = search_article(article, provider=RawStaticProvider([candidate]))
    assert not RepostRecord.objects.filter(article=article).exists()

    owned_channel.is_active = True
    owned_channel.save(update_fields=["is_active"])
    output = StringIO()
    with patch("apps.sources.services.fuzz.ratio", return_value=85.0):
        call_command("reclassify_automatic_repost_sites", stdout=output)

    saved = RepostRecord.objects.get(article=article)
    assert saved.content_relation == ContentRelation.OWNED
    assert saved.owned_channel is not None
    assert saved.owned_channel.code == "EASTMONEY"
    saved_candidate = SearchRunCandidate.objects.get(search_run=run)
    assert saved_candidate.content_relation == ContentRelation.OWNED
    assert "OWNED_CHANNEL:EASTMONEY=1" in output.getvalue()


@pytest.mark.django_db
def test_automatic_rule_change_does_not_downgrade_retained_external_repost(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    candidate = SearchCandidate(title=article.title, url="https://stable.example.com/article")
    search_article(article, provider=StaticProvider([candidate]))
    record = RepostRecord.objects.get(article=article)
    assert record.content_relation == ContentRelation.REPOST

    OwnedChannel.objects.create(
        code="STABLE_SITE",
        name="后续配置的自有站点",
        channel_type="OFFICIAL_WEBSITE",
        match_rules={"domains": ["stable.example.com"]},
    )
    second_run = search_article(article, provider=StaticProvider([candidate]))

    record.refresh_from_db()
    assert record.content_relation == ContentRelation.REPOST
    assert second_run.owned_count == 1
    assert SearchRunCandidate.objects.get(search_run=second_run).content_relation == ContentRelation.OWNED


@pytest.mark.django_db
def test_workspace_query_and_trend_exclude_owned_but_keep_sticky_reposts(source: Source, operator: User) -> None:
    first = create_source_article(source=source, operator=operator)
    second = Article.objects.create(
        title="第二篇可统计文章",
        normalized_title="第二篇可统计文章",
        published_date=first.published_date,
        published_at=first.published_at,
        channel_code="RIGHTS",
        channel_name="证券维权",
        section_code="COMPANY",
        section_name="涉事公司报道",
        monitoring_status=ArticleMonitoringStatus.COMPLETED,
        created_by=operator,
    )
    now = timezone.now()
    for article, suffix, relation in [
        (first, "one", ContentRelation.REPOST),
        (first, "owned", ContentRelation.OWNED),
        (second, "two", ContentRelation.REPOST),
    ]:
        url = f"https://trend.example.com/{suffix}"
        RepostRecord.objects.create(
            article=article,
            site_name="趋势站",
            site_domain="trend.example.com",
            raw_url=url,
            canonical_url=url,
            canonical_url_hash=sha256(url.encode("utf-8")).hexdigest(),
            original_url=url,
            normalized_url=url,
            normalized_url_hash=sha256(url.encode("utf-8")).hexdigest(),
            repost_title=article.title,
            result_title=article.title,
            first_discovered_at=now,
            first_found_at=now,
            last_checked_at=now,
            last_seen_at=now,
            availability_status="REMOVED" if suffix == "one" else "AVAILABLE",
            content_relation=relation,
            data_source="TEST_ONLY",
        )
    queryset = filtered_repost_articles(
        RepostQueryFilters(channel="RIGHTS"), as_of=timezone.now() + timedelta(seconds=1)
    )
    assert list(queryset.values_list("id", flat=True)) == [second.id]
    all_articles = filtered_repost_articles(RepostQueryFilters(), as_of=timezone.now() + timedelta(seconds=1))
    assert result_summary(all_articles, as_of=timezone.now() + timedelta(seconds=1))["repost_url_count"] == 2
    assert repost_trend(all_articles, as_of=timezone.now() + timedelta(seconds=1))[0]["repost_url_count"] == 2


@pytest.mark.django_db
@override_settings(SEARCH_CANDIDATE_MIN_SIMILARITY=0)
def test_result_workspace_is_paginated_and_export_is_read_only(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    now = timezone.now()
    record = RepostRecord.objects.create(
        article=article,
        site_name="外部站点",
        site_domain="external.example.com",
        raw_url="https://external.example.com/a",
        canonical_url="https://external.example.com/a",
        canonical_url_hash="c" * 64,
        original_url="https://external.example.com/a",
        normalized_url="https://external.example.com/a",
        normalized_url_hash="c" * 64,
        repost_title=article.title,
        result_title=article.title,
        first_discovered_at=now,
        first_found_at=now,
        last_checked_at=now,
        last_seen_at=now,
        data_source="TEST_ONLY",
    )
    client = APIClient()
    client.force_authenticate(operator)
    response = client.get(f"/api/v1/repost-monitor/results?published_from={article.published_date}&page=1&page_size=20")
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["pagination"] == {"page": 1, "page_size": 20, "total": 1}
    assert data["monitoring_status_counts"][ArticleMonitoringStatus.ACTIVE] == 1
    assert data["summary"]["repost_url_count"] == 1
    assert data["trend"][0]["repost_url_count"] == 1
    detail = client.get(f"/api/v1/articles/{article.id}/discovered-publications?as_of={data['as_of']}")
    assert detail.status_code == 200
    assert detail.json()["data"]["reposts"][0]["id"] == record.id
    run = search_article(
        article,
        provider=StaticProvider([SearchCandidate(title="不达阈值候选", url="https://external.example.com/candidate")]),
    )
    detail = client.get(f"/api/v1/articles/{article.id}/discovered-publications?as_of={data['as_of']}")
    assert detail.json()["data"]["candidates"][0]["search_run_id"] == run.id
    assert detail.json()["data"]["candidates"][0]["canonical_url"] == "https://external.example.com/candidate"
    with patch("apps.sources.services.configured_search_provider") as configured_provider:
        export = client.get(f"/api/v1/repost-monitor/export.xlsx?as_of={data['as_of']}")
    assert export.status_code == 200
    configured_provider.assert_not_called()


@pytest.mark.django_db
def test_provider_error_keeps_article_active_and_existing_reposts(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    search_article(
        article,
        provider=StaticProvider([SearchCandidate(title=article.title, url="https://example.com/repost")]),
    )
    run = search_article(article, provider=ErrorProvider())
    article.refresh_from_db()
    assert run.status == SearchRunStatus.ERROR
    assert run.error_code == "SEARCH_PROVIDER_RATE_LIMITED"
    assert article.monitoring_status == ArticleMonitoringStatus.ACTIVE
    assert RepostRecord.objects.filter(article=article).count() == 1


@pytest.mark.django_db
@override_settings(SEARCH_PROVIDER="", BRAVE_SEARCH_API_KEY="")
def test_unconfigured_provider_records_explicit_error_without_fake_results(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    SearchProviderConfiguration.objects.update(enabled=False)
    run = search_article(article)
    assert run.status == SearchRunStatus.ERROR
    assert run.error_code == "SEARCH_PROVIDER_NOT_CONFIGURED"
    assert RepostRecord.objects.filter(article=article).count() == 0


@pytest.mark.django_db
def test_repeated_discovery_updates_last_seen_without_changing_first_found(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    provider = StaticProvider([SearchCandidate(title=article.title, url="https://example.com/sticky")])
    search_article(article, provider=provider)
    record = RepostRecord.objects.get(article=article)
    first_found = record.first_found_at
    search_article(article, provider=provider)
    record.refresh_from_db()
    assert RepostRecord.objects.filter(article=article).count() == 1
    assert record.first_found_at == first_found
    assert record.last_seen_at is not None and first_found is not None and record.last_seen_at >= first_found


@pytest.mark.django_db
def test_retention_only_purges_after_article_retention(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    search_article(
        article,
        provider=StaticProvider([SearchCandidate(title=article.title, url="https://example.com/repost")]),
    )
    article.retention_until = timezone.now() + timedelta(days=1)
    article.save(update_fields=["retention_until"])
    assert purge_expired_source_articles()["articles_deleted"] == 0
    article.retention_until = timezone.now() - timedelta(seconds=1)
    article.save(update_fields=["retention_until"])
    assert purge_expired_source_articles()["articles_deleted"] == 1
    assert not Article.objects.filter(id=article.id).exists()


@pytest.mark.django_db
def test_dynamic_export_keeps_removed_links_and_creates_safe_site_sheets(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    now = timezone.now()
    for index, url in enumerate(("https://news.example.com/a", "https://news.example.com/b"), start=1):
        RepostRecord.objects.create(
            article=article,
            platform=None,
            site_name="危险:/网站名称很长很长很长很长很长很长很长很长",
            site_domain="news.example.com",
            raw_url=url,
            canonical_url=url,
            canonical_url_hash=str(index) * 64,
            original_url=url,
            normalized_url=url,
            normalized_url_hash=str(index) * 64,
            repost_title=f"{article.title} {index}",
            result_title=f"{article.title} {index}",
            similarity_score=95,
            first_discovered_at=now,
            first_found_at=now,
            last_checked_at=now,
            last_seen_at=now,
            availability_status="REMOVED" if index == 1 else "AVAILABLE",
            data_source="TEST_ONLY",
        )
    content, _ = build_global_repost_workbook(GlobalExportFilters(article_id=article.id))
    workbook = load_workbook(BytesIO(content))
    assert workbook.sheetnames[0] == "数据"
    assert workbook.sheetnames[-3:-1] == ["总统计", "统计汇总"]
    assert len(workbook.sheetnames) == 5
    assert len(workbook.sheetnames[1]) <= 31
    assert not INVALID_SHEET_CHARACTERS_FOR_TEST.search(workbook.sheetnames[1])
    assert "共2条" in workbook["数据"]["I2"].value
    assert workbook["数据"]["I2"].hyperlink.target == "https://news.example.com/a"
    assert workbook["数据"]["E2"].value == "监测中"
    assert workbook["数据"]["A2"].number_format == EXCEL_DATETIME_FORMAT
    assert workbook["数据"]["F2"].number_format == EXCEL_DATETIME_FORMAT
    site_sheet = workbook[workbook.sheetnames[1]]
    assert site_sheet.max_row == 3
    assert site_sheet["J2"].value == "已删除"
    assert site_sheet["H2"].number_format == EXCEL_DATETIME_FORMAT
    assert site_sheet["B2"].hyperlink.target == "https://news.example.com/a"
    assert workbook["总统计"]["F2"].value == "√"
    assert workbook["统计汇总"]["C2"].number_format == EXCEL_DATETIME_FORMAT


@pytest.mark.django_db
def test_export_filters_use_the_same_as_of_for_repost_existence(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    export_as_of = timezone.now()
    record = RepostRecord.objects.create(
        article=article,
        site_name="并发网站",
        site_domain="concurrent.example.com",
        raw_url="https://concurrent.example.com/repost",
        canonical_url="https://concurrent.example.com/repost",
        canonical_url_hash="a" * 64,
        original_url="https://concurrent.example.com/repost",
        normalized_url="https://concurrent.example.com/repost",
        normalized_url_hash="a" * 64,
        repost_title=article.title,
        result_title=article.title,
        first_discovered_at=export_as_of,
        first_found_at=export_as_of,
        last_checked_at=export_as_of,
        last_seen_at=export_as_of,
        data_source="TEST_ONLY",
    )
    RepostRecord.objects.filter(id=record.id).update(created_at=export_as_of + timedelta(seconds=1))

    assert (
        not filtered_articles(GlobalExportFilters(has_repost=True), as_of=export_as_of).filter(id=article.id).exists()
    )
    assert filtered_articles(GlobalExportFilters(has_repost=False), as_of=export_as_of).filter(id=article.id).exists()
    assert (
        not filtered_articles(GlobalExportFilters(site_domain="concurrent.example.com"), as_of=export_as_of)
        .filter(id=article.id)
        .exists()
    )


@pytest.mark.django_db
def test_export_link_totals_are_unique_per_article_and_url(source: Source, operator: User) -> None:
    first = create_source_article(source=source, operator=operator)
    second = Article.objects.create(
        title="第二篇文章",
        normalized_title="第二篇文章",
        published_date=timezone.localdate(),
        published_at=timezone.now(),
        monitoring_status=ArticleMonitoringStatus.COMPLETED,
        created_by=operator,
    )
    now = timezone.now()
    shared_url = "https://shared.example.com/repost"
    for article in (first, second):
        RepostRecord.objects.create(
            article=article,
            site_name="共享网站",
            site_domain="shared.example.com",
            raw_url=shared_url,
            canonical_url=shared_url,
            canonical_url_hash="b" * 64,
            original_url=shared_url,
            normalized_url=shared_url,
            normalized_url_hash="b" * 64,
            repost_title=article.title,
            result_title=article.title,
            first_discovered_at=now,
            first_found_at=now,
            last_checked_at=now,
            last_seen_at=now,
            data_source="TEST_ONLY",
        )

    content, _ = build_global_repost_workbook(GlobalExportFilters(site_domain="shared.example.com"))
    workbook = load_workbook(BytesIO(content))
    summary = workbook["统计汇总"]
    summary_values = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(2, 7)}
    assert summary_values["转载链接总数"] == 2
    assert summary.cell(9, 4).value == 2


@pytest.mark.django_db
@override_settings(
    ARTICLE_SEARCH_SCHEDULE_MINUTES=(0, 15, 30, 60, 120, 240, 480),
    ARTICLE_SEARCH_REPEAT_MINUTES=720,
)
def test_search_schedule_repeats_every_twelve_hours_after_eight_hours(source: Source, operator: User) -> None:
    article = create_source_article(source=source, operator=operator)
    started = timezone.now().replace(second=0, microsecond=0)
    article.monitor_started_at = started
    article.monitor_until = started + timedelta(days=7)

    assert _next_search_time(article, started + timedelta(hours=8)) == started + timedelta(hours=20)
    assert _next_search_time(article, started + timedelta(hours=19)) == started + timedelta(hours=20)
    assert _next_search_time(article, started + timedelta(hours=20)) == started + timedelta(hours=32)


INVALID_SHEET_CHARACTERS_FOR_TEST = __import__("re").compile(r"[:\\/?*\[\]]")


def test_sheet_name_sanitization_handles_collisions() -> None:
    used = {"数据", "总统计", "统计汇总"}
    first = sanitize_excel_sheet_name("网站:/名称", used)
    second = sanitize_excel_sheet_name("网站__名称", used)
    assert first == "网站__名称"
    assert second == "网站__名称_2"
