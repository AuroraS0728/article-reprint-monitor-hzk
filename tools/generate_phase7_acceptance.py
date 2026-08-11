"""Generate the Phase 7 acceptance-result DOCX from the reviewed Markdown facts."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "10_项目验收单_V1.0_第7轮测试结果.docx"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
RED = "9B1C1C"
GREEN = "1F6B45"
WIDTH_DXA = 9360


def set_run_font(
    run: object, *, size: float, bold: bool = False, color: str = "000000"
) -> None:
    run.font.name = "Calibri"  # type: ignore[attr-defined]
    run.font.size = Pt(size)  # type: ignore[attr-defined]
    run.bold = bold  # type: ignore[attr-defined]
    run.font.color.rgb = RGBColor.from_string(color)  # type: ignore[attr-defined]
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")  # type: ignore[attr-defined]
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")  # type: ignore[attr-defined]


def shade(cell: object, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()  # type: ignore[attr-defined]
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell: object) -> None:
    tc = cell._tc  # type: ignore[attr-defined]
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), "80" if side in {"top", "bottom"} else "120")
        node.set(qn("w:type"), "dxa")


def format_table(table: object, widths: list[int], header: bool = True) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT  # type: ignore[attr-defined]
    table.autofit = False  # type: ignore[attr-defined]
    table_pr = table._tbl.tblPr  # type: ignore[attr-defined]
    table_width = table_pr.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    table_pr.append(indent)
    for row_index, row in enumerate(table.rows):  # type: ignore[attr-defined]
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    set_run_font(run, size=9.5, bold=row_index == 0 and header)
            if row_index == 0 and header:
                shade(cell, LIGHT_BLUE)
        if row_index % 2 == 1 and header:
            for cell in row.cells:
                shade(cell, "FAFBFC")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    paragraph.paragraph_format.space_after = Pt(8 if level == 1 else 6)
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color=BLUE)


def add_text(
    doc: Document, text: str, *, bold: bool = False, color: str = "000000"
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    run = paragraph.add_run(text)
    set_run_font(run, size=11, bold=bold, color=color)


def add_table(
    doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers, strict=True):
        cell.text = text
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values, strict=True):
            cell.text = text
    format_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("文章转载监测系统 V1.0 | 第 7 轮测试结果 | 2026-08-10")
        set_run_font(run, size=8.5, color="666666")


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("文章转载监测系统"), size=24, bold=True, color="0B2545")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(
        subtitle.add_run("项目验收单 - 第 7 轮测试结果"), size=15, bold=True, color=BLUE
    )

    add_table(
        doc,
        ["项目", "内容"],
        [
            ["验收版本", "V1.0 / 第 7 轮"],
            ["测试日期", "2026-08-10"],
            [
                "测试环境",
                "Windows 11 本机；Docker Desktop；Python 3.12.13 Alpine；MySQL 8.4.4；Redis 7.4.2；Node 22.14.0",
            ],
            ["代码基线", "ade95c3 - fix(security): harden runtime audit environment"],
            ["验收结论", "不允许上线"],
        ],
        [2700, 6660],
    )
    add_text(
        doc,
        "本记录是《项目验收单 V1.0》的第 7 轮实测补充，不修改原始待签署验收模板。",
        color="555555",
    )

    add_heading(doc, "1. 已通过的自动化检查")
    add_table(
        doc,
        ["类别", "实际命令/范围", "结果"],
        [
            ["后端测试", "pytest -q", "65 passed（7.84s）"],
            ["格式", "black --check .", "111 个文件符合格式"],
            ["静态检查", "ruff check .", "通过"],
            ["类型检查", "mypy apps config", "96 个源文件无问题"],
            ["Django 部署检查", "manage.py check --deploy", "0 issues"],
            ["Python 依赖", "pip-audit --local", "未发现已知漏洞"],
            ["前端 lint", "pnpm lint", "通过"],
            ["前端生产依赖", "pnpm audit --prod --audit-level high", "未发现已知漏洞"],
            ["TypeScript/生产构建", "pnpm build", "通过"],
            ["密钥扫描", "scan_tracked_secrets.py", "未发现常见凭据形式"],
            ["数据库备份", "create_database_backup()", "成功；ID 8；54,915 bytes"],
        ],
        [1900, 4100, 3360],
    )
    add_text(
        doc,
        "后端测试覆盖登录限制、角色/越权、Excel 上传与公式注入、平台请求 SSRF/DNS 防护、检测状态 1/0/—、批次调度、日报/周报、人工补录、审计、日志脱敏、清理与备份逻辑。该覆盖不等同于浏览器端到端测试或真实平台检出率。",
    )

    add_heading(doc, "2. Docker 镜像漏洞检查")
    add_text(
        doc,
        "最终镜像 article-reprint-monitor-web:latest（digest 83a6f0ad142c）由 Docker Scout 扫描：0 Critical、1 High、0 Medium、0 Low，镜像约 86 MB、145 个包。",
    )
    warning = doc.add_table(rows=1, cols=1)
    warning.cell(0, 0).text = (
        "唯一 High：mariadb 11.8.8-r0 的 CVE-2025-13699。Scout 的 Alpine 数据显示“<= 11.8.8-r0，未提供修复版”；"
        "MariaDB 官方 CVE 清单显示该 CVE 已在 Community Server 11.8.4 修复。镜像内 mysqldump 已实测为 11.8.8，"
        "故按上游版本判定为扫描器包元数据误报/已修复版本。不得忽略扫描规则；后续镜像更新时必须复核。"
    )
    format_table(warning, [9360], header=False)
    shade(warning.cell(0, 0), "FFF4E5")
    for run in warning.cell(0, 0).paragraphs[0].runs:
        set_run_font(run, size=10, color=RED)
    add_text(
        doc,
        "运行镜像使用非 root appuser 且具有可写 home 目录；Python 依赖审计和数据库备份均在该权限模型下实际通过。MySQL 与 Redis 未发布宿主机端口；web:8000 和 frontend:5173 是本机开发端口，不得作为生产暴露策略。",
    )

    add_heading(doc, "3. 尚未通过或未执行的验收项")
    add_table(
        doc,
        ["项目", "状态", "原因/后续动作"],
        [
            [
                "真实转载检出率 >= 95%",
                "未执行",
                "尚无经人工确认的正样本集合；不得把历史 Excel 空白当负样本。",
            ],
            ["误报率 <= 2%", "未执行", "尚无经人工确认的负样本集合。"],
            [
                "五平台正式适配器真实验证",
                "未完成",
                "URL 探测证据不能替代完整搜索/解析适配器验收。",
            ],
            ["浏览器 E2E", "未执行", "本机浏览器自动化运行时未可用，需补测。"],
            ["SMTP 真实投递", "未执行", "仅 mock/逻辑测试，未配置真实授权信息。"],
            [
                "数据库恢复演练",
                "未执行",
                "逻辑备份已成功，仍须在隔离数据库完成恢复校验。",
            ],
            [
                "24 小时稳定性、P95、告警、回滚",
                "未执行",
                "需独立持续运行和运维测试环境。",
            ],
            ["生产网络/Nginx/域名/TLS", "未执行", "本轮不部署正式生产。"],
        ],
        [2700, 1500, 5160],
    )

    add_heading(doc, "4. 最终结论")
    conclusion = doc.add_paragraph()
    conclusion.paragraph_format.space_after = Pt(8)
    set_run_font(conclusion.add_run("不允许上线。"), size=13, bold=True, color=RED)
    set_run_font(
        conclusion.add_run(
            " 当前无未判定 Critical 漏洞；唯一 High 扫描告警已依据官方修复版本和实际版本完成证据化判定。"
            "但生产验收仍缺少真实平台 KPI、恢复演练、真实 SMTP、持续运行及生产网络验证。"
            "完成这些项目并签署原始验收单后，才可批准上线。"
        ),
        size=11,
    )
    add_text(
        doc,
        "后续动作：在每次基础镜像或 MariaDB 客户端更新后复跑 Docker Scout 与 pip-audit；完成真实样本集、恢复演练和生产环境验收后，创建新的验收版本，不覆盖本记录。",
        bold=True,
        color=GREEN,
    )
    add_footer(doc)
    doc.core_properties.title = "文章转载监测系统 V1.0 第 7 轮测试结果"
    doc.core_properties.author = "文章转载监测系统"
    doc.core_properties.subject = "第 7 轮测试与安全验收补充记录"
    doc.core_properties.comments = "Automated evidence only; not production approval."
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
